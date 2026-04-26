import os
from typing import Dict, List, Tuple

import httpx
from .GenerateCommands_L3 import GenerateCommands
from .GenerateProcedure.GenerateProcedure import GenerateProcedure
from .Models import TcRequest,GeneratedTestProcedure
from .Logging.Logger import logger
import polars as pl
import copy 
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
# make dir in C:\\Users\\Public\\temp\\ if not exists
dir_name = "C:\\Users\\Public\\temp"
os.makedirs(dir_name, exist_ok=True)

logger.info("Creating Instance of GenerateCommands and GenerateProcedure")
gc = GenerateCommands()
gp = GenerateProcedure(gc.tc_details_df,gc.r)
gc.gui_status_update_raw(summary="Application", status="Ready...")
go_lang_api = "http://127.0.0.1:11000/"

class ApiHandler:

    @classmethod
    async def generate_test_procedure_document(cls,tc_request: TcRequest):
        # Create a new document
        try:
            doc = Document()
            doc.add_heading(f'Payload Test Procedures', level=1)
            cfg_numbers = gc.get_config_numbers_list_from_str(tc_request.configs_str)
            tc_request =await  gc.pre_process_request(tc_request)
            if len(cfg_numbers) == 0:
                raise ValueError("No Configurations Found in the Request")
            progress_percentage = 100 / len(cfg_numbers)
            tc_request.config_numbers = cfg_numbers
            commands_df = await gc.generate_commands_for_test_procedure_document(tc_request)
            progress = 0
            for cfg in cfg_numbers:
                tc_request.configs_str = str(cfg)
                tc_request = await gc.pre_process_request(tc_request)
                tc_request.line_number = 1
                cfg_commands_df = commands_df.filter(pl.col("CFG_NO").is_in(tc_request.config_numbers))
                if cfg_commands_df.height > 0:
                    commands = cfg_commands_df["TC"].to_list()
                    tc_request.commands = commands
                    procedure = await gp.generate_procedure(tc_request)
                    tc_request.commands = []
                    tc_request = await  gc.generate_boa_commands(tc_request)
                    procedure += await gp.generate_boa_procedure(tc_request)
                    if len(procedure)>1:
                        procedure =  procedure + gp.end_of_procedure()
                        doc.add_heading(f'Test Procedure for Config Number: {cfg}', level=2)
                        # Add your generated text for this iteration
                        # Add as a paragraph in monospace font
                        p = doc.add_paragraph()
                        run = p.add_run(procedure)
                        run.font.name = 'Consolas'  # or 'Courier New'
                        run.font.size = Pt(11)
                        # For compatibility with Word, set the font in underlying XML
                        r = run._element
                        r.rPr.rFonts.set(qn('w:eastAsia'), 'Consolas')
                        #doc.add_paragraph(procedure)
                        # Add extra space if needed
                        doc.add_paragraph('')  
                    # write the procedure to word document


                progress += progress_percentage
                gc.update_status(f"Completed for CFG:{cfg}",progress,"Procedures Generation In Progress...")

            doc.save(f"{dir_name}{os.sep}procedures.docx")
            gc.update_status("All Procedures Generated",100,"Procedures Generation Completed")
            return "OK"
        except Exception as e:
            logger.error(f"Error occurred while generating test procedure document: {e}")
            raise Exception(f"Error occurred while generating test procedure document: {e}")

    @classmethod
    async def generate_test_procedure_turn_off_document(cls,tc_request: TcRequest):
        # Create a new document
        try:
            doc = Document()
            doc.add_heading(f'Payload Test Procedures', level=1)
            cfg_numbers = gc.get_config_numbers_list_from_str(tc_request.configs_str)
            tc_request =await  gc.pre_process_request(tc_request)
            if len(cfg_numbers) == 0:
                raise ValueError("No Configurations Found in the Request")
            progress_percentage = 100 / len(cfg_numbers)
            tc_request.config_numbers = cfg_numbers
            commands_df = await gc.generate_commands_for_test_procedure_off_document(tc_request)
            progress = 0
            for cfg in cfg_numbers:
                tc_request.configs_str = str(cfg)
                tc_request = await gc.pre_process_request(tc_request)
                cfg_commands_df = commands_df.filter(pl.col("CFG_NO").is_in(tc_request.config_numbers))
                if cfg_commands_df.height > 0:
                    commands = cfg_commands_df["TC"].to_list()
                    off_commands = []
                    for cmd in commands:
                        off_cmd = gc.on_to_off_cmd_df.filter(pl.col("ON_CMD") == cmd)
                        if off_cmd.height > 0:
                            off_commands.append(off_cmd["OFF_CMD"].to_list()[0])
                        else:
                            off_commands.append(cmd[:-2] + "OFF;")  # Assuming the command ends with ON;
                    tc_request.commands = off_commands
                    procedure = await gp.generate_procedure(tc_request)
                    tc_request.commands = []
                    # tc_request = await  gc.generate_boa_commands(tc_request)
                    # procedure += await gp.generate_boa_procedure(tc_request)
                    if len(procedure)>1:
                        procedure =  procedure + gp.end_of_procedure()
                        doc.add_heading(f'Test Procedure for Config Number: {cfg}', level=2)
                        # Add your generated text for this iteration
                        # Add as a paragraph in monospace font
                        p = doc.add_paragraph()
                        run = p.add_run(procedure)
                        run.font.name = 'Consolas'  # or 'Courier New'
                        run.font.size = Pt(11)
                        # For compatibility with Word, set the font in underlying XML
                        r = run._element
                        r.rPr.rFonts.set(qn('w:eastAsia'), 'Consolas')
                        #doc.add_paragraph(procedure)
                        # Add extra space if needed
                        doc.add_paragraph('')  
                    # write the procedure to word document


                progress += progress_percentage
                gc.update_status(f"Completed for CFG:{cfg}",progress,"Procedures Generation In Progress...")

            doc.save(f"{dir_name}{os.sep}procedures.docx")
            gc.update_status("All Procedures Generated",100,"Procedures Generation Completed")
            return "OK"
        except Exception as e:
            logger.error(f"Error occurred while generating test procedure document: {e}")
            raise Exception(f"Error occurred while generating test procedure document: {e}")

    
    @classmethod
    async def set_configs_on(cls,tc_request: TcRequest)->GeneratedTestProcedure:
        tc_request = await gc.pre_process_request(tc_request)
        pre_procedure,line_number = gc.get_additional_procedure_at_the_start_of_procedure(1)
        tc_req =await  gc.generate_cfg_based_on_commands(tc_request)
        if len(pre_procedure) >0:
            tc_req.line_number = line_number
        procedure = await gp.generate_procedure(tc_req)
        if len(procedure) >0:
            tc_req.line_number += 2
        tc_req.commands = []
        if not tc_req.change_only_switches:
            tc_req = await  gc.generate_boa_commands(tc_req)
            boa_procedure = await gp.generate_boa_procedure(tc_req)
            procedure += boa_procedure
        if len(procedure)>1:
            if tc_req.change_only_switches:
                procedure = procedure + gp.end_of_procedure()
            else:
                procedure = pre_procedure + procedure + gp.end_of_procedure()
        else:
            gc.update_status("Noting to Generate",100,"TM Matched")
        generated_test_procedure = GeneratedTestProcedure(file_name=tc_request.file_name)
        generated_test_procedure.file_content = procedure
        return generated_test_procedure
    
    @classmethod
    async def set_configs_off(cls,tc_request: TcRequest)->GeneratedTestProcedure:
        p = ""
        tc_request = await gc.pre_process_request(tc_request)
        tc_request.line_number = 1
        gp.line_number = 1
        tc_request.boa_column_name = "max"
        tc_boa_req = await  gc.generate_boa_commands(tc_request)
        boa_procedure = await gp.generate_boa_procedure(tc_boa_req)
        tc_req = await  gc.generate_cfg_based_off_commands(tc_request)
        tc_req.line_number = gp.line_number
        procedure = await gp.generate_procedure(tc_req)
        if len(procedure)>1:
            p += boa_procedure+procedure+ gp.end_of_procedure()
        else:
            gc.update_status("Noting to Generate",100,"TM Matched")
        
        generated_test_procedure = GeneratedTestProcedure(file_name=tc_request.file_name)
        generated_test_procedure.file_content = p
        return generated_test_procedure
    
    @classmethod
    async def set_all_payload_off(cls)-> GeneratedTestProcedure:
        p = ""
        tc_request = TcRequest()
        tc_request = await gc.pre_process_request(tc_request)
        tc_req =await gc.generate_all_rf_systems_off_commands(tc_request)
        procedure = await gp.generate_procedure(tc_req)
        if len(procedure)>1:
            p += procedure+ gp.end_of_procedure()
        else:
            gc.update_status("Noting to Generate",100,"TM Matched")
        generated_test_procedure = GeneratedTestProcedure(file_name=tc_request.file_name)
        generated_test_procedure.file_content = p
        return generated_test_procedure
    
        
    
    @classmethod
    async def generate_data_commands(cls,tc_request: TcRequest)->GeneratedTestProcedure:
        procedure = ""
        tc_request = await gc.pre_process_request(tc_request)
        boa_procedure = await gp.generate_boa_procedure(tc_request)
        gc.update_status("",48,"Procedure Generated")
        procedure += boa_procedure
        if len(procedure)>1:
            procedure += gp.end_of_procedure()
        else:
            gc.update_status("Noting to Generate",100,"TM Matched")
        generated_test_procedure = GeneratedTestProcedure(file_name=tc_request.file_name)
        generated_test_procedure.file_content = procedure
        return generated_test_procedure
    
    @classmethod
    async def manual_commands(cls,tc_request: TcRequest)->GeneratedTestProcedure:
        tc_request.manual_commanding_mode = True
        tc_request = await gc.pre_process_request(tc_request)
        pre_procedure,line_number = gc.get_additional_procedure_at_the_start_of_procedure(1)
        if len(pre_procedure) >0:
            tc_request.line_number = line_number
        tc_req = await  gc.generate_manual_command(tc_request)
        procedure = await gp.generate_procedure(tc_req)
        if len(procedure)>1:
            procedure = pre_procedure + procedure + gp.end_of_procedure()
        else:
            gc.update_status("Noting to Generate",100,"TM Matched")
        generated_test_procedure = GeneratedTestProcedure(file_name=tc_request.file_name)
        generated_test_procedure.file_content = procedure
        return generated_test_procedure

    @classmethod
    async def generate_boa_cal_procedure(cls,tc_request: TcRequest)->List[str]:
        tc_request = await gc.pre_process_request(tc_request)
        tc_request.line_number = 1
        tp = await gp.generate_boa_cal_procedure(tc_request)
        if len(tp)>1:
            tp = tp + gp.end_of_procedure()
        generated_test_procedure = GeneratedTestProcedure(file_name=tc_request.file_name,file_content=tp)
        return generated_test_procedure
    
    @classmethod
    async def get_test_parameter_names(cls)->List[str]:
        return gc.get_test_parameter_names()
    
    @classmethod
    async def get_cfg_boa_col_names(cls)->List[str]:
        return gc.get_cfg_boa_col_names()
    
    @classmethod
    async def get_tc_file_names(cls)->List[str]:
        return gc.get_tc_file_names()
    
    @classmethod
    async def get_commands_group_names(cls)->List[str]:
        return gc.get_commands_group_names()
    
    @classmethod
    async def get_manual_commands(self)->Dict[str, List[str]]:
        return gc.get_manual_commands()
    
    @classmethod
    async def get_data_commands_list(cls)->List[str]:
        return gc.get_data_commands_list()
    
    @classmethod
    async def get_data_command_values(cls,data_cmd: str)->Dict[str,str]:
        return gc.get_data_commands_values(data_cmd)
    
    
    @classmethod
    async def execute_test_procedure(cls,request:GeneratedTestProcedure):
        payload = {
            "proc_name": request.file_name,
            "procedure": request.file_content
        }
        async with httpx.AsyncClient() as client:
            response = await client.post(go_lang_api+"umacs_tc_transfer_file_and_trigger_execution", json=payload)
            if response.status_code == 200:
                gc.update_status(f"Waiting For File:{request.file_name} Execution...",10,"File Triggered in UMACS")
                return "OK"
            else:
                gc.update_status(response.text,0,f"Failed to execute file:{request.file_name}")
                raise Exception(response.text)
        return "OK"
    
    @classmethod
    async def trigger_file_execution(cls,request:GeneratedTestProcedure):
        payload = {
            "proc_name": request.file_name,
            "procedure": ""
        }
        async with httpx.AsyncClient() as client:
            response = await client.post(go_lang_api+"umacs_tc_trigger_file_execution", json=payload)
            if response.status_code == 200:
                gc.gui_status_update_raw(status =f"Waiting For File:{request.file_name} Execution...",progress=10,summary="File Triggered in UMACS")
                return "OK"
            else:
                gc.gui_status_update_raw(status=response.text,progress=0,summary=f"Failed to execute file:{request.file_name}")
                #raise HTTPException(status_code=response.status_code, detail=response.text)
        return "OK"

